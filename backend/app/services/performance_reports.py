"""Native EMS report rendering, adapted from the author's MAEstro report workflow.

Only the template/counter grouping and document ergonomics are reused. This module
does not connect to MAE, consume corporate cookies, or embed screenshots.
"""
from collections import defaultdict
from datetime import datetime, timezone
from io import BytesIO
from threading import Lock

from app.services.performance import performance_service


RENDER_LOCK = Lock()


def plot_series(axis, series, granularity):
    """Draw observed runs separately: isolated samples must remain visible."""
    points = sorted(series['points'], key=lambda p: p['epoch'])
    segments = []
    for point in points:
        if not segments or point['epoch'] - segments[-1][-1]['epoch'] > granularity * 1.5:
            segments.append([])
        segments[-1].append(point)
    color = None
    for index, segment in enumerate(segments):
        line, = axis.plot(
            [datetime.fromisoformat(p['timestamp']) for p in segment],
            [p['value'] for p in segment],
            label=series['object_id'] if index == 0 else '_nolegend_',
            color=color, linewidth=1.8, marker='o' if len(segment) < 30 else None,
            markersize=3, zorder=3,
            drawstyle='steps-post' if series['counter_id'] == 'core.nf.availability' else 'default',
        )
        color = line.get_color()
    return max(0, len(segments) - 1)


def configure_axis(axis, series_list):
    if all(s['counter_id'] == 'core.nf.availability' for s in series_list):
        # Keep 0 and 100 away from the frame. Do not clip the 100% plateau.
        axis.set_ylim(-5, 105)
        axis.set_yticks([0, 20, 40, 60, 80, 100])
    axis.margins(x=.03)


def build_report(request, user):
    # One reference instant for all relative queries; SQLite history is the source.
    end = datetime.now(timezone.utc)
    results = []
    for item in request.queries:
        payload = item.model_dump()
        payload["end"] = payload.get("end") or end
        result = performance_service.query(payload, user)
        results.append((item.name, result))
    if sum(len(result["series"]) for _, result in results) > 60:
        raise ValueError("El informe admite hasta 60 series. Reduzca las consultas u objetos.")
    with RENDER_LOCK:
        return render(request.title, user.username, results, end)


def render(title, author, results, generated_at):
    from docx import Document
    from docx.shared import Inches, Pt
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    import matplotlib.dates as mdates

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(title, 0)
    doc.add_paragraph(f"EMS educativo · {author} · {generated_at:%Y-%m-%d %H:%M:%S} UTC")
    doc.add_paragraph("Informe generado desde muestras históricas del EMS, sin capturas manuales. Ausencia de muestras no equivale a cero. Las estadísticas siguientes se calculan sobre los puntos agregados de cada consulta, no sobre toda la población de paquetes.")
    for name, result in results:
        doc.add_heading(name, 1)
        doc.add_paragraph(f"Testbed: {result['testbed_id']} | Escenario: {result['scenario_id']}\nPeriodo UTC: {result['start']} — {result['end']}\nResolución: {result['granularity_seconds']} s | Muestras: {result['sample_count']}")
        groups = defaultdict(list)
        for series in result["series"]:
            groups[series["counter_id"]].append(series)
        if not groups:
            doc.add_paragraph("Sin muestras para esta consulta en el periodo solicitado.")
        if result.get("missing_series"):
            doc.add_paragraph(f"Advertencia: {len(result['missing_series'])} combinaciones objeto/contador no tienen muestras.")
        for _, series_list in groups.items():
            heading = series_list[0]["label"].split(" · ", 1)[-1]
            doc.add_heading(heading, 2)
            fig = Figure(figsize=(8.8, 3.1), layout="constrained")
            canvas = FigureCanvasAgg(fig)
            axis = fig.subplots()
            gaps = sum(plot_series(axis, series, result['granularity_seconds']) for series in series_list)
            configure_axis(axis, series_list)
            axis.set_ylabel(series_list[0]["unit"])
            axis.xaxis.set_major_formatter(mdates.DateFormatter("%d/%m %H:%M", tz=timezone.utc))
            axis.tick_params(axis="x", labelsize=7)
            axis.grid(alpha=.2)
            axis.legend(fontsize=7)
            image = BytesIO()
            canvas.print_png(image)
            image.seek(0)
            doc.add_picture(image, width=Inches(6.3))
            if gaps:
                doc.add_paragraph(f"Sin muestras: {gaps} interrupciones entre segmentos de las series. Los espacios sin línea son periodos sin observaciones, no disponibilidad del 0 % ni del 100 %. Los puntos aislados representan una sola muestra.")
            table = doc.add_table(rows=1, cols=5)
            table.style = "Light Shading Accent 1"
            for cell, text in zip(table.rows[0].cells, ["Objeto", "Último", "Mínimo", "Máximo", "Promedio / tipo"]):
                cell.text = text
            for series in series_list:
                values = [p["value"] for p in series["points"]]
                counter = series.get("kind") == "counter"
                stats = [series["object_id"], f"{values[-1]:.3f}", f"{min(values):.3f}", f"{max(values):.3f}", "Acumulado; no se promedia" if counter else f"{sum(values)/len(values):.3f}"]
                for cell, text in zip(table.add_row().cells, stats):
                    cell.text = text
            sources = sorted({f"{s['source']} ({s['quality']})" for s in series_list})
            doc.add_paragraph("Fuente: " + ", ".join(sources) + ". Unidad: " + series_list[0]["unit"])
        for note in result.get("notes", []):
            doc.add_paragraph(note)
    doc.add_heading("Alcance de la medición", 1)
    doc.add_paragraph("Las métricas nativas dependen de la versión/configuración del emulador. Una NF compartida entre escenarios mide el proceso completo. PLMN, S-NSSAI, DNN y causas se conservan como series distintas. Los acumulados pueden reiniciarse; las tasas derivadas no atraviesan reinicios ni pausas de recolección superiores a 120 segundos. Este informe no demuestra conformidad 3GPP ni aislamiento de slices.")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
